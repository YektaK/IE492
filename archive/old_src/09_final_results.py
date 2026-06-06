# -*- coding: utf-8 -*-
"""
09_final_results.py
Consolidates every analysis output into a single workbook
D:/IE492/output/final/Sultanbeyli_Final_Results.xlsx

Sheets (in order):
  0. README
  1. KPI_Summary
  2. Selected_8
  3. AHP_Weights
  4. TOPSIS_Top20
  5. Sensitivity
  6. Coverage_per_Mahalle
  7. Coverage_KPI
  8. Inputs_Candidates
  9. Inputs_Existing
 10. Inputs_Mahalle_Risk
 11. Inputs_Mahalle_Pop
 12. Mu_Aday
 13. Mu_Mevcut
 14. Centroids
 15. Criteria_Matrix
 16. AHP_Robustness
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule

ROOT = Path("D:/IE492")
DATA = ROOT / "output" / "data"
RES = ROOT / "output" / "results"
FINAL = ROOT / "output" / "final"
FINAL.mkdir(parents=True, exist_ok=True)

OUT = FINAL / "Sultanbeyli_Final_Results.xlsx"

# ---------------------------------------------------------------------------
# Load all source dataframes
# ---------------------------------------------------------------------------
adaylar = pd.read_excel(DATA / "adaylar.xlsx")
mevcut = pd.read_excel(DATA / "mevcut_12.xlsx")
nufus = pd.read_excel(DATA / "mahalle_nufus.xlsx")
risk = pd.read_excel(DATA / "mahalle_risk.xlsx")

ahp = pd.read_excel(RES / "ahp_weights.xlsx")
topsis = pd.read_excel(RES / "topsis_sonuclar.xlsx")
sel = pd.read_excel(RES / "selection_Baseline_hard.xlsx")
sens = pd.read_excel(RES / "selection_sensitivity.xlsx")
cov = pd.read_excel(RES / "coverage_score_detailed.xlsx")
kpi = pd.read_excel(RES / "coverage_score_kpi.xlsx")
centroids = pd.read_excel(RES / "mahalle_centroids.xlsx")
mu_aday = pd.read_excel(RES / "mu_aday.xlsx")
mu_mevcut = pd.read_excel(RES / "mu_mevcut.xlsx")
criteria = pd.read_excel(RES / "criteria_matrix.xlsx")

# ---------------------------------------------------------------------------
# Build cover-sheet KPIs
# ---------------------------------------------------------------------------
dist_17_baseline = 76.5
dist_17_proposed = 82.4
urban_15_baseline = 86.7
urban_15_proposed = 93.3
mean_mu_base = 0.658
mean_mu_prop = 0.716
selected_ids = sel["S_No"].tolist()
mahalle_of = dict(zip(adaylar["S_No"], adaylar["Mahalle"]))
sel_by_mah = {}
for sid in selected_ids:
    m = mahalle_of.get(sid, "?")
    sel_by_mah[m] = sel_by_mah.get(m, 0) + 1

kpi_rows = [
    ("KPI", "Value"),
    ("Candidate pool (AYDES)", len(adaylar)),
    ("Mahalle (district)", len(risk)),
    ("Existing containers", len(mevcut)),
    ("New containers to add", 8),
    ("Final total network", 20),
    ("AHP CR (Baseline)", 0.0545),
    ("FCM sigma (m)", 800),
    ("IP status", "Optimal"),
    ("IP objective Z", 935.621365),
    ("Coverage uplift (continuous mu)", "+96.6 %"),
    ("Mahalle coverage eşik 0.50 (17)", f"{dist_17_baseline}% -> {dist_17_proposed}% (+5.9 pp)"),
    ("Mahalle coverage eşik 0.50 (15 kentsel)", f"{urban_15_baseline}% -> {urban_15_proposed}% (+6.6 pp)"),
    ("Mean mu (17 mahalle)", f"{mean_mu_base:.3f} -> {mean_mu_prop:.3f}"),
    ("Hamidiye mu (rescue)", f"0.473 -> 0.992 (S59..S62)"),
    ("Sensitivity scenarios (AHP x IP)", "3 x 2 = 6 — all converge"),
    ("Selected S_No", ", ".join("S" + str(s) for s in selected_ids)),
    ("Sites by mahalle", ", ".join(f"{m}:{n}" for m, n in sel_by_mah.items())),
]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
HDR_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
SUB_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN = Side(border_style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def write_df(ws, df, header=True, header_fill=HDR_FILL, header_font=WHITE):
    if header:
        ws.append(list(df.columns))
        for c in ws[ws.max_row]:
            c.fill = header_fill
            c.font = header_font
            c.alignment = CENTER
            c.border = BOX
    for row in dataframe_to_rows(df, index=False, header=False):
        ws.append(row)
    # auto width
    for col in ws.columns:
        m = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(m + 2, 38)

def add_title(ws, text, ncols, fill=HDR_FILL, font=WHITE):
    ws.append([text])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=ncols)
    c = ws.cell(row=ws.max_row, column=1)
    c.fill = fill
    c.font = font
    c.alignment = CENTER

# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------
wb = Workbook()

# ---- 0. README --------------------------------------------------------------
ws = wb.active
ws.title = "README"
add_title(ws, "SULTANBEYLI KONTEYNER OPTIMIZASYONU - FINAL RESULTS", 1, fill=HDR_FILL, font=WHITE)
readme = [
    ("", ""),
    ("Ders", "IE 492 - Endüstri Mühendisliği Bitirme Projesi"),
    ("Danisman", "Dr. Ahmet Yekta Kayman"),
    ("Ekip", "Elif Keleş, Zümra Sancaklı, Doğa Yardemir, Semanur Aydın"),
    ("Tarih", "Ocak 2026"),
    ("", ""),
    ("Problem", "Mevcut 12 AFIS konteynerine ek olarak 8 yeni konteyner konumu secimi (Valilik karari: hedef 20)."),
    ("Yontem", "AHP -> TOPSIS -> FCM (Gaussian, sigma=800m) -> 0-1 IP (PuLP/CBC)."),
    ("Aday havuzu", f"{len(adaylar)} AYDES kayitli kamu alani."),
    ("Mahalle sayisi", f"{len(risk)} (ilce geneli)."),
    ("", ""),
    ("Secilen 8", ", ".join("S" + str(s) for s in selected_ids)),
    ("Obj Z", "935.62  (Optimal, 6/6 senaryo)"),
    ("Kapsama artisi", "+96.6 % (surekli FCM mu)"),
    ("Mahalle kapsama (eşik 0.50, 17 m.)", "76.5% -> 82.4%"),
    ("Mahalle kapsama (eşik 0.50, 15 kentsel)", "86.7% -> 93.3%"),
    ("", ""),
    ("SHEET LISTESI", ""),
    ("1. KPI_Summary", "Tum basari olculeri tek tabloda"),
    ("2. Selected_8", "8 yeni konteyner (S_No, ad, mahalle, koordinat, TOPSIS skoru)"),
    ("3. AHP_Weights", "3 senaryo agirlik vektorleri + CR"),
    ("4. TOPSIS_Top20", "En yuksek 20 TOPSIS skoru (her senaryo)"),
    ("5. Sensitivity", "3 AHP x 2 IP modu (hepsi ayni sonuc)"),
    ("6. Coverage_per_Mahalle", "17 mahalle icin mu bazli kapsama"),
    ("7. Coverage_KPI", "Mahalle seviyesi kapsama KPI"),
    ("8. Inputs_Candidates", "140 aday (C1..C4)"),
    ("9. Inputs_Existing", "12 mevcut konteyner"),
    ("10. Inputs_Mahalle_Risk", "IBB Tablo 5-2 risk"),
    ("11. Inputs_Mahalle_Pop", "2024 nufus verisi"),
    ("12. Mu_Aday", "140x17 FCM mu matrisi (aday)"),
    ("13. Mu_Mevcut", "12x17 FCM mu matrisi (mevcut)"),
    ("14. Centroids", "Mahalle merkez koordinatlari"),
    ("15. Criteria_Matrix", "AHP ikili karsilastirma matrisi"),
    ("16. AHP_Robustness", "Mikro perturbasyon testi (rho = 0.99942)"),
]
for r in readme:
    ws.append(list(r))
ws.column_dimensions["A"].width = 36
ws.column_dimensions["B"].width = 90
for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for c in r:
        c.alignment = LEFT
        c.border = BOX
        if c.row == 1:
            c.fill = HDR_FILL
            c.font = WHITE
        elif c.value and isinstance(c.value, str) and c.value.isupper() and " " not in c.value.strip():
            c.fill = SUB_FILL
            c.font = BOLD

# ---- 1. KPI_Summary ---------------------------------------------------------
ws = wb.create_sheet("KPI_Summary")
write_df(ws, pd.DataFrame(kpi_rows[1:], columns=kpi_rows[0]))

# ---- 2. Selected_8 -----------------------------------------------------------
ws = wb.create_sheet("Selected_8")
# selection file already has all the C1..C4, TOPSIS, lat/lon we need
write_df(ws, sel)

# ---- 3. AHP_Weights ----------------------------------------------------------
ws = wb.create_sheet("AHP_Weights")
add_title(ws, "AHP Weights - 3 Scenarios", 6)
hdr = ["Scenario", "w_C1_risk", "w_C2_logistics", "w_C3_gap", "w_C4_nightpop", "CR"]
ws.append(hdr)
for c in ws[ws.max_row]:
    c.fill = HDR_FILL; c.font = WHITE; c.alignment = CENTER
ws.append(["Baseline", 0.541, 0.254, 0.117, 0.088, 0.0545])
ws.append(["DamageFocused", 0.644, 0.194, 0.093, 0.069, 0.0615])
ws.append(["InfrastructureFocused", 0.406, 0.406, 0.105, 0.082, 0.0395])
ws.append([])
ws.append(["Note", "All CR < 0.10 - tutarli", "", "", "", ""])
ws.column_dimensions["A"].width = 24
for col in "BCDEF":
    ws.column_dimensions[col].width = 16

# number-format the weight cells
last_data_row = ws.max_row - 1
for r in ws.iter_rows(min_row=2, max_row=last_data_row, min_col=2, max_col=6):
    for c in r:
        c.number_format = "0.0000"

# ---- 4. TOPSIS_Top20 ---------------------------------------------------------
ws = wb.create_sheet("TOPSIS_Top20")
top20 = topsis.head(20)
write_df(ws, top20)

# ---- 5. Sensitivity ----------------------------------------------------------
ws = wb.create_sheet("Sensitivity")
write_df(ws, sens)

# ---- 6. Coverage_per_Mahalle -------------------------------------------------
ws = wb.create_sheet("Coverage_per_Mahalle")
write_df(ws, cov)

# ---- 7. Coverage_KPI ---------------------------------------------------------
ws = wb.create_sheet("Coverage_KPI")
write_df(ws, kpi)

# ---- 8. Inputs_Candidates ----------------------------------------------------
ws = wb.create_sheet("Inputs_Candidates")
write_df(ws, adaylar)

# ---- 9. Inputs_Existing ------------------------------------------------------
ws = wb.create_sheet("Inputs_Existing")
write_df(ws, mevcut)

# ---- 10. Inputs_Mahalle_Risk -------------------------------------------------
ws = wb.create_sheet("Inputs_Mahalle_Risk")
write_df(ws, risk)

# ---- 11. Inputs_Mahalle_Pop --------------------------------------------------
ws = wb.create_sheet("Inputs_Mahalle_Pop")
write_df(ws, nufus)

# ---- 12. Mu_Aday -------------------------------------------------------------
ws = wb.create_sheet("Mu_Aday")
write_df(ws, mu_aday)

# ---- 13. Mu_Mevcut -----------------------------------------------------------
ws = wb.create_sheet("Mu_Mevcut")
write_df(ws, mu_mevcut)

# ---- 14. Centroids -----------------------------------------------------------
ws = wb.create_sheet("Centroids")
write_df(ws, centroids)

# ---- 15. Criteria_Matrix -----------------------------------------------------
ws = wb.create_sheet("Criteria_Matrix")
write_df(ws, criteria)

# ---- 16. AHP_Robustness -----------------------------------------------------
ws = wb.create_sheet("AHP_Robustness")
add_title(ws, "AHP Weight Robustness - intra-scenario test", 2)
rows = [
    ("Test", "Value"),
    ("Test type", "Micro-perturbation within Baseline scenario"),
    ("Weight set A (ours)", "[0.541, 0.254, 0.117, 0.088]"),
    ("Weight set B (alt)", "[0.534, 0.265, 0.114, 0.087]"),
    ("Spearman rank correlation (rho)", 0.99942),
    ("p-value", 1.35e-204),
    ("Mean |CC_A - CC_B|", 0.00521),
    ("Max |CC_A - CC_B|", 0.01152),
    ("Top-10 order", "Identical (S2, S62, S4, S61, S7, S81, S5, S60, S59, S84)"),
    ("0-1 IP final selection", "Unchanged (S4, S59, S60, S61, S62, S83, S88, S141)"),
    ("Interpretation", "Solution is AHP-robust: FCM coverage constraint is binding, not AHP weighting"),
]
for r in rows:
    ws.append(list(r))
ws.column_dimensions["A"].width = 36
ws.column_dimensions["B"].width = 72
# format numeric cells
for r_idx in (5, 6, 7, 8):
    cell = ws.cell(row=r_idx + 1, column=2)  # +1 because title row
    if r_idx == 5:  # rho
        cell.number_format = "0.00000"
    elif r_idx == 6:  # p-value
        cell.number_format = "0.00E+00"
    else:
        cell.number_format = "0.00000"

# ---------------------------------------------------------------------------
# Conditional formatting on Coverage_per_Mahalle
# ---------------------------------------------------------------------------
import openpyxl
ws_cov = wb["Coverage_per_Mahalle"]
for col_idx in range(1, ws_cov.max_column + 1):
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    if "mu" in str(ws_cov.cell(row=1, column=col_idx).value).lower() or \
       "coverage" in str(ws_cov.cell(row=1, column=col_idx).value).lower():
        rng = f"{col_letter}2:{col_letter}{ws_cov.max_row}"
        ws_cov.conditional_formatting.add(
            rng,
            ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B"))

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
wb.save(OUT)
print(f"[OK] {OUT}  sheets={len(wb.sheetnames)}")
print(f"     Sheets: {wb.sheetnames}")
